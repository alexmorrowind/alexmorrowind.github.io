from __future__ import annotations

import os
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
PPTX_PATH = OUT / "B1Pay_Octo_Partner_Proposal.pptx"
ONE_PAGER_PATH = OUT / "B1Pay_Octo_One_Pager.pdf"

FONT_DIR = Path("/System/Library/Fonts/Supplemental")
ARIAL_REG = FONT_DIR / "Arial Unicode.ttf"
ARIAL_BOLD = FONT_DIR / "Arial Bold.ttf"
if ARIAL_REG.exists():
    pdfmetrics.registerFont(TTFont("B1Arial", str(ARIAL_REG)))
if ARIAL_BOLD.exists():
    pdfmetrics.registerFont(TTFont("B1Arial-Bold", str(ARIAL_BOLD)))

W = Inches(13.333)
H = Inches(7.5)

NAVY = RGBColor(17, 42, 77)
BLUE = RGBColor(18, 100, 214)
CYAN = RGBColor(44, 177, 213)
INK = RGBColor(32, 45, 62)
MUTED = RGBColor(86, 108, 135)
PALE = RGBColor(242, 247, 252)
PALE_BLUE = RGBColor(233, 242, 255)
PALE_GREEN = RGBColor(237, 248, 241)
PALE_ORANGE = RGBColor(255, 244, 223)
PALE_RED = RGBColor(255, 241, 237)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(183, 203, 226)


def add_full_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_footer(slide, idx, label="B1Pay × Octo | Концепция для переговоров"):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(7.12), Inches(12.38), Inches(0.008))
    line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background()
    add_text(slide, 0.52, 7.18, 9.0, 0.18, label, 8, MUTED)
    add_text(slide, 12.35, 7.16, 0.45, 0.2, str(idx), 9, MUTED, align=PP_ALIGN.RIGHT)


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP, margin=0.04, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb


def add_rich_text(slide, x, y, w, h, runs, size=18, color=INK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.04); tf.margin_right = Inches(0.04)
    p = tf.paragraphs[0]; p.alignment = align
    for item in runs:
        run = p.add_run(); run.text = item.get("text", "")
        run.font.name = item.get("font", "Aptos")
        run.font.size = Pt(item.get("size", size))
        run.font.bold = item.get("bold", False)
        run.font.color.rgb = item.get("color", color)
    return tb


def add_title(slide, title, subtitle=None, section=None):
    if section:
        add_text(slide, 0.55, 0.28, 3.5, 0.24, section.upper(), 9, BLUE, True)
    add_text(slide, 0.52, 0.55, 12.1, 0.58, title, 27, NAVY, True)
    if subtitle:
        add_text(slide, 0.55, 1.17, 11.9, 0.42, subtitle, 12, MUTED)


def box(slide, x, y, w, h, title, body, fill=PALE_BLUE, title_color=NAVY, body_color=MUTED, accent=None, title_size=14, body_size=10.5):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = LINE; sh.line.width = Pt(1)
    if accent:
        a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
        a.fill.solid(); a.fill.fore_color.rgb = accent; a.line.fill.background()
    add_text(slide, x + 0.18, y + 0.15, w - 0.32, 0.32, title, title_size, title_color, True)
    add_text(slide, x + 0.18, y + 0.55, w - 0.32, h - 0.65, body, body_size, body_color)
    return sh


def bullet_box(slide, x, y, w, h, title, bullets, fill=PALE_BLUE, accent=BLUE, title_size=14, body_size=11):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = LINE; sh.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.075), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    add_text(slide, x+0.18, y+0.15, w-0.3, 0.3, title, title_size, NAVY, True)
    tb = slide.shapes.add_textbox(Inches(x+0.18), Inches(y+0.53), Inches(w-0.33), Inches(h-0.65))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.02); tf.margin_right = Inches(0.02)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(body_size)
        p.font.color.rgb = INK
        p.space_after = Pt(5)
        p.bullet = True
    return sh


def connector(slide, x1, y1, x2, y2, color=LINE, width=2, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    if dashed:
        c.line.dash_style = 1
    return c


def add_brand(slide, dark=False):
    color = WHITE if dark else NAVY
    add_rich_text(slide, 0.55, 0.18, 2.0, 0.33, [
        {"text": "B1", "bold": True, "color": BLUE, "size": 18},
        {"text": "pay", "bold": True, "color": color, "size": 18},
    ])


def new_slide(prs, idx, title, subtitle=None, section=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_bg(slide, NAVY if dark else WHITE)
    if not dark:
        add_brand(slide)
        add_title(slide, title, subtitle, section)
        add_footer(slide, idx)
    return slide


def build_pptx():
    prs = Presentation()
    prs.slide_width = W; prs.slide_height = H
    prs.core_properties.title = "B1Pay × Octo — предложение о партнёрстве"
    prs.core_properties.subject = "Партнёрская модель, архитектура и пилот"
    prs.core_properties.author = "B1Pay"
    prs.core_properties.keywords = "B1Pay, Octo, Uzbekistan, fintech, P2P, MyID"

    # 1 — cover
    s = new_slide(prs, 1, "B1Pay × Octo", "Партнёрская модель цифровой финансовой платформы для Узбекистана", dark=True)
    add_text(s, 0.72, 1.52, 11.2, 0.9, "Сравнить → выбрать → подтвердить → получить сервис через лицензированного партнёра", 28, WHITE, True)
    add_text(s, 0.76, 2.75, 10.6, 0.76, "Каталог банков и финансовых продуктов сейчас. P2P и другие регулируемые операции — после юридического и технического согласования.", 16, RGBColor(214, 229, 246))
    box(s, 0.76, 4.35, 3.6, 1.25, "B1Pay", "UX, каталог, заявки, MyID/QR-сессии, аналитика", fill=RGBColor(31, 62, 104), title_color=WHITE, body_color=RGBColor(211, 227, 245), accent=CYAN)
    box(s, 4.87, 4.35, 3.6, 1.25, "Octo / банк", "Платёжный контур и регулируемая операция в согласованной роли", fill=RGBColor(31, 62, 104), title_color=WHITE, body_color=RGBColor(211, 227, 245), accent=BLUE)
    box(s, 8.98, 4.35, 3.6, 1.25, "Пользователь", "Единый интерфейс, прозрачные условия и статус заявки", fill=RGBColor(31, 62, 104), title_color=WHITE, body_color=RGBColor(211, 227, 245), accent=RGBColor(245, 181, 71))
    add_text(s, 0.78, 6.65, 10.5, 0.28, "Концепция для переговоров • 10 сентября 2026", 10, RGBColor(176, 198, 222))

    # 2 — ask
    s = new_slide(prs, 2, "Что предлагаем Octo", "Начать с управляемого пилота и расширять периметр только после approval gates", "01 | Резюме")
    bullet_box(s, 0.65, 1.8, 5.9, 3.9, "Пилот без лишнего риска", [
        "Каталог и заявки на кредиты, микрозаймы, карты и страхование.",
        "MyID в мобильном приложении; web показывает QR и не принимает биометрию.",
        "1–2 продукта, ограниченная группа пользователей, sandbox, мониторинг и reconciliation.",
        "P2P выключен feature flag до юридического, договорного и технического подтверждения.",
    ], fill=PALE_BLUE, accent=BLUE, body_size=12)
    bullet_box(s, 6.8, 1.8, 5.9, 3.9, "Что хотим согласовать", [
        "Разделить merchant/acquiring/e-commerce API и отдельный P2P/payment-service workstream.",
        "Допустимую роль B1Pay и перечень разрешённых Octo-сценариев.",
        "API, подписи, webhooks, лимиты, комиссии, settlement и dispute flow.",
        "KYC/AML RACI: кто принимает решение, хранит доказательства и отвечает за отчётность.",
        "План go/no-go для P2P/payout после успешного контролируемого пилота — без обещания лицензии.",
    ], fill=PALE_GREEN, accent=RGBColor(36, 150, 95), body_size=12)
    add_text(s, 0.7, 6.05, 11.7, 0.55, "Принцип: B1Pay не заявляет себя банком или платёжной организацией; регулируемая операция выполняется лицензированным партнёром в его контуре.", 14, NAVY, True, align=PP_ALIGN.CENTER)

    # 3 — current baseline
    s = new_slide(prs, 3, "От чего отталкиваемся", "В репозитории уже есть технический каркас; коммерческий и регулируемый контур нужно согласовать", "02 | Текущий задел")
    box(s, 0.65, 1.75, 3.75, 1.45, "Web-портал", "Банки, кредиты, калькулятор, карточки продуктов, SEO-страницы, QR-вход.", fill=PALE_BLUE, accent=BLUE)
    box(s, 4.55, 1.75, 3.75, 1.45, "Flutter Android", "Профиль, MyID-абстракция, QR confirmation, инвестиционный UI, статусы.", fill=PALE_BLUE, accent=CYAN)
    box(s, 8.45, 1.75, 3.9, 1.45, "Django API", "Общие пользователи, банки, заявки, инвестиции и интеграционные endpoints.", fill=PALE_GREEN, accent=RGBColor(36, 150, 95))
    connector(s, 4.4, 2.47, 4.55, 2.47, BLUE, 2); connector(s, 8.3, 2.47, 8.45, 2.47, BLUE, 2)
    bullet_box(s, 0.65, 3.75, 5.8, 2.2, "Уже предусмотрено", [
        "Поля/адаптеры для Octo: base URL, merchant/shop ID, secret и P2P endpoint.",
        "P2P-флаг по умолчанию false — безопасная база для пилота.",
        "MyID session/status и QR-вход с подтверждением на телефоне.",
    ], fill=PALE_GREEN, accent=RGBColor(36, 150, 95), body_size=11.3)
    bullet_box(s, 6.7, 3.75, 5.65, 2.2, "Что нельзя считать готовым", [
        "Наличие лицензии, production credentials или разрешения на P2P.",
        "Готовность банковских тарифов и предложений без partner approval.",
        "Юридическую квалификацию инвестиционного модуля.",
    ], fill=PALE_RED, accent=RGBColor(194, 85, 72), body_size=11.3)

    # 4 value
    s = new_slide(prs, 4, "Ценность платформы", "B1Pay сокращает путь от поиска финансового продукта до подтверждённой заявки", "03 | Продукт")
    stages = [("1", "Найти", "Каталог с фильтрами, источником и датой обновления", PALE_BLUE), ("2", "Понять", "Калькулятор, комиссии, риски и требования", PALE_ORANGE), ("3", "Подтвердить", "MyID на телефоне, QR для web, consent", PALE_GREEN), ("4", "Получить", "Статус от банка/МФО/Octo, договор у партнёра", PALE_RED)]
    xs = [0.62, 3.85, 7.08, 10.31]
    for i, (num, name, body, fill) in enumerate(stages):
        box(s, xs[i], 2.15, 2.45, 2.55, f"{num}  {name}", body, fill=fill, accent=[BLUE, RGBColor(245, 181, 71), RGBColor(36, 150, 95), RGBColor(194, 85, 72)][i], title_size=16, body_size=11)
        if i < 3: connector(s, xs[i] + 2.48, 3.42, xs[i+1] - 0.07, 3.42, BLUE, 2)
    add_text(s, 0.75, 5.25, 11.8, 0.62, "Прозрачность: каждая карточка показывает, кто предоставляет продукт, кто принимает решение, где находится договор и как обновлены условия.", 16, NAVY, True, align=PP_ALIGN.CENTER)

    # 5 roles
    s = new_slide(prs, 5, "Роли и границы ответственности", "Партнёрская модель делает owner-of-record и клиентский договор понятными", "04 | Operating model")
    role_data = [
        ("B1Pay", "UX, каталог, routing, consent, support L1", PALE_BLUE, BLUE),
        ("Octo / банк", "API, settlement, лимиты, fraud, платежный статус", PALE_ORANGE, RGBColor(245, 181, 71)),
        ("Банк / МФО", "Кредитное решение, договор, выдача и взыскание", PALE_GREEN, RGBColor(36, 150, 95)),
        ("Страховщик", "Андеррайтинг, полис, claims", PALE_RED, RGBColor(194, 85, 72)),
        ("MyID", "KYC session и результат в согласованном периметре", PALE_BLUE, CYAN),
        ("Пользователь", "Выбор, согласие, подтверждение, распоряжение", PALE_ORANGE, RGBColor(245, 181, 71)),
    ]
    coords = [(0.65, 1.8), (4.47, 1.8), (8.29, 1.8), (0.65, 4.1), (4.47, 4.1), (8.29, 4.1)]
    for (name, body, fill, accent), (x, y) in zip(role_data, coords):
        box(s, x, y, 3.45, 1.75, name, body, fill=fill, accent=accent, title_size=15, body_size=11)
    add_text(s, 0.72, 6.3, 11.7, 0.35, "B1Pay не принимает на себя роль регулируемого поставщика без отдельного основания, лицензии и договора.", 13, RGBColor(194, 85, 72), True, align=PP_ALIGN.CENTER)

    # 6 product map
    s = new_slide(prs, 6, "Продуктовая карта: сейчас и после согласования", "От независимой витрины к partner-led financial rails", "05 | Roadmap")
    box(s, 0.65, 1.75, 5.8, 4.35, "Сейчас — marketplace / information layer", "• Банки, кредиты, микрозаймы\n• Карты и депозиты\n• Страхование\n• Калькуляторы и сравнение\n• Заявка / lead routing\n• MyID + QR-вход\n• Инвестиционный каталог с раскрытием рисков", fill=PALE_BLUE, accent=BLUE, body_size=14)
    box(s, 6.85, 1.75, 5.75, 4.35, "Два разных workstream после согласования", "• Merchant/acquiring/e-commerce API — отдельный технический контур\n• P2P / card-to-card / payment-service — отдельная legal/licensing проверка\n• Webhooks и reconciliation\n• Ограниченные лимиты и пилотная группа\n• Статусы, возвраты, dispute/chargeback\n• Отдельный legal perimeter инвестиций", fill=PALE_GREEN, accent=RGBColor(36, 150, 95), body_size=13.2)
    add_text(s, 0.75, 6.35, 11.75, 0.3, "На сайте и в приложении можно заранее показать P2P как «В разработке», не создавая обещания доступности.", 13, MUTED, True, align=PP_ALIGN.CENTER)

    # 7 P2P flow
    s = new_slide(prs, 7, "P2P-перевод: целевой поток после согласования", "Final status приходит от Octo/банка; клиент не меняет его локально", "06 | P2P")
    steps = [("1", "Пользователь", "Выбирает перевод, сумму и получателя"), ("2", "B1Pay", "Проверяет сессию, consent, лимит и device risk"), ("3", "Octo / банк", "Signed API request + AML/fraud/limits"), ("4", "Webhook", "pending → success / failed / reversed"), ("5", "B1Pay", "Reconcile, receipt, support/dispute")]
    xs = [0.6, 3.05, 5.5, 7.95, 10.4]
    for i, (n, title, body) in enumerate(steps):
        box(s, xs[i], 2.15, 2.15, 2.6, f"{n}  {title}", body, fill=[PALE_BLUE, PALE_BLUE, PALE_ORANGE, PALE_GREEN, PALE_GREEN][i], accent=[BLUE, BLUE, RGBColor(245, 181, 71), RGBColor(36, 150, 95), RGBColor(36, 150, 95)][i], title_size=14, body_size=10.5)
        if i < 4: connector(s, xs[i]+2.18, 3.45, xs[i+1]-0.04, 3.45, BLUE, 2)
    bullet_box(s, 0.75, 5.35, 5.6, 1.12, "Обязательные controls", ["Idempotency key · webhook signature · TTL · reconciliation"], fill=PALE_RED, accent=RGBColor(194, 85, 72), body_size=11)
    bullet_box(s, 6.95, 5.35, 5.5, 1.12, "Go / no-go gate", ["Legal role · contract · sandbox · limits · AML/RACI · incident runbook"], fill=PALE_ORANGE, accent=RGBColor(245, 181, 71), body_size=11)

    # 8 KYC QR
    s = new_slide(prs, 8, "MyID + QR: безопасная связка web и телефона", "Web показывает только короткоживущую сессию; подтверждение выполняется на телефоне", "07 | Identity")
    box(s, 0.65, 1.8, 3.3, 3.65, "Web", "1. Создать session\n2. Получить QR с nonce/TTL\n3. Ждать signed completion\n\nНе хранит биометрию и не принимает фото документа.", fill=PALE_BLUE, accent=BLUE, body_size=12)
    box(s, 4.97, 1.8, 3.3, 3.65, "Mobile app", "1. Пользователь входит через MyID\n2. Сканирует/подтверждает QR\n3. Показывает понятное действие\n4. Отправляет approve одноразово", fill=PALE_GREEN, accent=RGBColor(36, 150, 95), body_size=12)
    box(s, 9.29, 1.8, 3.3, 3.65, "B1Pay backend", "TTL · nonce · device binding\nreplay prevention · audit\n\nПередаёт партнёру только нужный минимум и статус сессии.", fill=PALE_ORANGE, accent=RGBColor(245, 181, 71), body_size=12)
    connector(s, 3.98, 3.58, 4.95, 3.58, BLUE, 2); connector(s, 8.3, 3.58, 9.26, 3.58, BLUE, 2)
    add_text(s, 0.75, 6.05, 11.7, 0.54, "Consent должен быть отдельным, понятным и журналироваться: цель, объём, получатель, время, срок и возможность отзыва.", 14, NAVY, True, align=PP_ALIGN.CENTER)

    # 9 architecture
    s = new_slide(prs, 9, "Системная схема", "Общий control plane для web, Android и партнёрских API", "08 | Architecture")
    # top clients
    box(s, 0.65, 1.55, 2.55, 1.0, "Web", "catalog · QR · applications", PALE_BLUE, accent=BLUE, body_size=10.5)
    box(s, 3.45, 1.55, 2.55, 1.0, "Flutter", "MyID · QR · status", PALE_BLUE, accent=CYAN, body_size=10.5)
    box(s, 6.25, 1.55, 2.55, 1.0, "Admin", "offers · cases · audit", PALE_ORANGE, accent=RGBColor(245, 181, 71), body_size=10.5)
    box(s, 9.05, 1.55, 3.55, 1.0, "Partner portal", "SLA · leads · reports", PALE_GREEN, accent=RGBColor(36, 150, 95), body_size=10.5)
    # control plane
    box(s, 1.15, 3.0, 11.1, 1.45, "B1Pay API + Integration Gateway", "Auth/session · catalog · consent ledger · application routing · payment state machine · webhook processor · reconciliation · risk/AML queue · audit/observability", fill=RGBColor(235, 242, 253), accent=BLUE, title_size=17, body_size=12)
    # providers
    box(s, 0.65, 5.25, 2.55, 1.05, "MyID", "KYC session/status", PALE_RED, accent=RGBColor(194, 85, 72), body_size=10.5)
    box(s, 3.45, 5.25, 2.55, 1.05, "Octo / bank", "P2P · payout · settlement", PALE_ORANGE, accent=RGBColor(245, 181, 71), body_size=10.5)
    box(s, 6.25, 5.25, 2.55, 1.05, "Bank / MFO", "credit decision", PALE_GREEN, accent=RGBColor(36, 150, 95), body_size=10.5)
    box(s, 9.05, 5.25, 3.55, 1.05, "Insurer / investment perimeter", "policy or separately governed module", PALE_RED, accent=RGBColor(194, 85, 72), body_size=10.5)
    for x in [1.92, 4.72, 7.52, 10.82]: connector(s, x, 2.58, x, 2.99, BLUE, 1.5)
    for x in [1.92, 4.72, 7.52, 10.82]: connector(s, x, 4.47, x, 5.22, BLUE, 1.5)

    # 10 compliance
    s = new_slide(prs, 10, "Compliance и безопасность — обязательная часть продукта", "Пилот должен быть ограничен, наблюдаем и готов к аудиту", "09 | Risk")
    bullet_box(s, 0.65, 1.75, 3.85, 4.35, "KYC / AML", [
        "Owner-of-record у лицензированного партнёра.",
        "Риск-уровни, санкции/PEP, лимиты и manual review.",
        "Письменный RACI и escalation path.",
        "Журнал согласий и доказательств.",
    ], fill=PALE_BLUE, accent=BLUE, body_size=11.5)
    bullet_box(s, 4.75, 1.75, 3.85, 4.35, "Security", [
        "Secrets только в server-side secret manager.",
        "HMAC/mTLS/IP allowlist для partner API.",
        "No raw PAN/CVV, passport photos or biometrics in logs.",
        "RBAC + MFA, monitoring, incident runbook.",
    ], fill=PALE_GREEN, accent=RGBColor(36, 150, 95), body_size=11.5)
    bullet_box(s, 8.85, 1.75, 3.8, 4.35, "Consumer protection", [
        "Источник, дата обновления и sponsored disclosure.",
        "Решение по кредиту/страхованию принимает партнёр.",
        "Комиссии, риски, возврат и complaint path понятны.",
        "Никаких обещаний гарантированного дохода.",
    ], fill=PALE_RED, accent=RGBColor(194, 85, 72), body_size=11.5)

    # 11 revenue
    s = new_slide(prs, 11, "Модель доходов и стимулы", "Монетизация должна быть прозрачной и не менять независимый рейтинг скрыто", "10 | Economics")
    revs = [
        ("CPL / CPA", "за квалифицированную заявку или одобренный продукт", PALE_BLUE, BLUE),
        ("Revenue share", "за операцию, если это допустимо договором", PALE_GREEN, RGBColor(36, 150, 95)),
        ("Sponsored", "платное размещение с явной маркировкой", PALE_ORANGE, RGBColor(245, 181, 71)),
        ("SaaS / white-label", "кабинет и API для партнёров", PALE_RED, RGBColor(194, 85, 72)),
    ]
    for i, (name, body, fill, accent) in enumerate(revs):
        box(s, 0.68 + i*3.16, 2.1, 2.75, 2.35, name, body, fill=fill, accent=accent, title_size=15, body_size=11.5)
    add_text(s, 0.8, 5.15, 11.5, 0.35, "Пользователь должен видеть, когда карточка оплачена, и всегда видеть официальный источник и дату обновления.", 14, NAVY, True, align=PP_ALIGN.CENTER)
    bullet_box(s, 1.2, 5.75, 5.05, 0.9, "Экономическая воронка", ["view → click → application → KYC → approval → operation"], fill=PALE_BLUE, accent=BLUE, body_size=10)
    bullet_box(s, 7.0, 5.75, 5.05, 0.9, "Контроль", ["CPL/CPA, fraud rate, SLA, dispute rate, revenue reconciliation"], fill=PALE_GREEN, accent=RGBColor(36, 150, 95), body_size=10)

    # 12 roadmap
    s = new_slide(prs, 12, "Дорожная карта запуска", "Порядок: сначала legal/discovery, затем каталог и пилот, потом регулируемые rails", "11 | Delivery")
    roadmap = [
        ("0", "Legal & discovery", "2–4 недели\nроль B1Pay, договор, RACI, sandbox", PALE_RED, RGBColor(194, 85, 72)),
        ("1", "Catalog & leads", "4–6 недель\noffers, consent, MyID/QR, status", PALE_BLUE, BLUE),
        ("2", "Controlled pilot", "6–10 недель\n1–2 партнёра, limits, monitoring", PALE_GREEN, RGBColor(36, 150, 95)),
        ("3", "Expansion", "после approval\nP2P/payout, more partners, white-label", PALE_ORANGE, RGBColor(245, 181, 71)),
    ]
    for i, (n, title, body, fill, accent) in enumerate(roadmap):
        x = 0.6 + i*3.16
        box(s, x, 2.0, 2.72, 2.75, f"{n}  {title}", body, fill=fill, accent=accent, title_size=14, body_size=11.5)
        if i < 3: connector(s, x+2.75, 3.35, x+3.1, 3.35, BLUE, 2)
    add_text(s, 0.8, 5.35, 11.5, 0.5, "P2P feature flag остаётся выключенным до прохождения юридического и технического approval gate.", 15, RGBColor(194, 85, 72), True, align=PP_ALIGN.CENTER)
    add_text(s, 0.8, 5.95, 11.5, 0.36, "Веб и приложение развиваются из одной API-модели, поэтому partner status и consent одинаковы на обеих платформах.", 12, MUTED, align=PP_ALIGN.CENTER)

    # 13 request
    s = new_slide(prs, 13, "Что просим у Octo на первой встрече", "Конкретный список решений делает пилот измеримым", "12 | Partner ask")
    bullet_box(s, 0.65, 1.75, 5.85, 4.55, "Коммерция и право", [
        "Какие merchant/acquiring/e-commerce API доступны в Octo?",
        "Какие P2P/card-to-card/payment-service сценарии допустимы отдельно?",
        "Какая юридическая роль нужна B1Pay и какие документы обязательны?",
        "Кто заключает договор с пользователем и кто owner-of-record по AML/KYC?",
        "Комиссии, settlement, SLA, chargeback/refund и support escalation.",
    ], fill=PALE_ORANGE, accent=RGBColor(245, 181, 71), body_size=11.4)
    bullet_box(s, 6.8, 1.75, 5.85, 4.55, "Техника и пилот", [
        "Sandbox, API docs, signing, webhook events and error codes.",
        "Лимиты, idempotency, reconciliation и тестовые реквизиты.",
        "Security questionnaire, IP allowlist, certificates, monitoring.",
        "Пилотная группа, продукт, срок, KPI и go/no-go критерии.",
    ], fill=PALE_BLUE, accent=BLUE, body_size=11.4)
    add_text(s, 0.8, 6.45, 11.5, 0.35, "Предложение: совместный discovery workshop на 90 минут и затем короткий sandbox plan.", 14, NAVY, True, align=PP_ALIGN.CENTER)

    # 14 success metrics
    s = new_slide(prs, 14, "Как поймём, что пилот работает", "Результат — не только транзакции, но и управляемый риск, качество данных и понятный UX", "13 | KPI")
    metrics = [
        ("Конверсия", "view → click → application → KYC → decision"),
        ("Надёжность", "success rate, pending time, webhook delivery, reversals"),
        ("Риск", "fraud rate, AML cases, complaint rate, dispute rate"),
        ("Экономика", "CPL/CPA, CAC, revenue per lead, reconciliation"),
    ]
    for i, (name, body) in enumerate(metrics):
        box(s, 0.8 + (i%2)*6.05, 1.95 + (i//2)*1.75, 5.45, 1.35, name, body, fill=[PALE_BLUE, PALE_GREEN, PALE_RED, PALE_ORANGE][i], accent=[BLUE, RGBColor(36, 150, 95), RGBColor(194, 85, 72), RGBColor(245, 181, 71)][i], title_size=15, body_size=12)
    add_text(s, 0.75, 5.7, 11.8, 0.75, "Главная метрика: пользователь получает прозрачный путь и правильный статус, а партнёр — качественную заявку/операцию в контролируемом контуре.", 17, NAVY, True, align=PP_ALIGN.CENTER)

    # 15 closing
    s = new_slide(prs, 15, "Следующий шаг", "B1Pay готовит UX и технический контур; Octo помогает определить допустимый regulated perimeter", "14 | Decision")
    add_text(s, 0.9, 1.8, 11.5, 0.55, "Предлагаем совместно подтвердить:", 22, NAVY, True, align=PP_ALIGN.CENTER)
    bullet_box(s, 2.0, 2.75, 9.3, 2.25, "Decision pack", [
        "роль B1Pay + список разрешённых продуктов",
        "sandbox/API + security/AML requirements",
        "пилотный продукт, группа, лимиты и KPI",
        "go/no-go критерии для P2P/payout",
    ], fill=PALE_BLUE, accent=BLUE, body_size=15)
    add_text(s, 1.0, 5.65, 11.3, 0.5, "Спасибо. Контакт для workshop: команда B1Pay", 19, BLUE, True, align=PP_ALIGN.CENTER)
    add_text(s, 1.0, 6.25, 11.3, 0.3, "Приложения: business_model.md • architecture.md • platform_architecture.svg", 11, MUTED, align=PP_ALIGN.CENTER)

    prs.save(PPTX_PATH)
    return PPTX_PATH


def build_one_pager():
    styles = getSampleStyleSheet()
    title = ParagraphStyle("Title", parent=styles["Title"], fontName="B1Arial-Bold", fontSize=20, leading=24, textColor=colors.HexColor("#112A4D"), alignment=TA_LEFT, spaceAfter=6)
    sub = ParagraphStyle("Sub", parent=styles["Normal"], fontName="B1Arial", fontSize=9.5, leading=13, textColor=colors.HexColor("#526984"), spaceAfter=9)
    h = ParagraphStyle("H", parent=styles["Heading2"], fontName="B1Arial-Bold", fontSize=12.5, leading=15, textColor=colors.HexColor("#1264D6"), spaceBefore=8, spaceAfter=4)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName="B1Arial", fontSize=9.2, leading=12.2, textColor=colors.HexColor("#202D3E"), spaceAfter=3)
    small = ParagraphStyle("Small", parent=body, fontSize=8.2, leading=10.5, textColor=colors.HexColor("#526984"))
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=10, firstLineIndent=-6, bulletIndent=0, spaceAfter=2)
    doc = SimpleDocTemplate(str(ONE_PAGER_PATH), pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=13*mm, bottomMargin=13*mm)
    story = []
    story.append(Paragraph("B1Pay × Octo — предложение о партнёрстве", title))
    story.append(Paragraph("Цифровая финансовая витрина и будущий partner-led платёжный слой для Узбекистана. Версия для переговоров, 10 сентября 2026.", sub))
    story.append(Paragraph("Идея", h))
    story.append(Paragraph("B1Pay объединяет каталог банковских продуктов, кредиты/микрозаймы, страхование, MyID + QR-вход и отдельный инвестиционный модуль. Пользователь сравнивает условия и направляет заявку лицензированному партнёру. P2P/payout подключается только после юридического, договорного, AML и технического согласования с Octo/банком.", body))
    story.append(Paragraph("Ценность для Octo / банка", h))
    vals = [
        "Структурированные лиды и единая API-точка для web и Android.",
        "Контролируемый пилот: sandbox, лимиты, SLA, webhooks, reconciliation.",
        "MyID на телефоне, QR в web, минимизация KYC/payment данных.",
        "Воронка и метрики: просмотр → заявка → KYC → решение → операция.",
    ]
    for x in vals: story.append(Paragraph("• " + x, bullet))
    story.append(Paragraph("Модель ответственности", h))
    table_data = [
        [Paragraph("B1Pay", small), Paragraph("UX, каталог, consent, routing, support L1", small)],
        [Paragraph("Octo / банк", small), Paragraph("API, settlement, лимиты, fraud и regulated operation в разрешённой роли", small)],
        [Paragraph("Банк/МФО/страховщик", small), Paragraph("Кредитное решение, договор, полис, выдача и claims", small)],
        [Paragraph("MyID", small), Paragraph("KYC session/status в согласованном контуре", small)],
    ]
    t = Table(table_data, colWidths=[38*mm, 135*mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#F2F7FC")),
        ("BOX", (0,0), (-1,-1), 0.6, colors.HexColor("#B7CBE2")),
        ("INNERGRID", (0,0), (-1,-1), 0.35, colors.HexColor("#D9E5F1")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 6), ("RIGHTPADDING", (0,0), (-1,-1), 6),
        ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story.append(t)
    story.append(Paragraph("Пилот", h))
    story.append(Paragraph("Фаза 0: legal/discovery → Фаза 1: каталог и лиды → Фаза 2: controlled pilot → Фаза 3: P2P/payout после approval gate. На первом этапе P2P остаётся выключенным feature flag и не обещается пользователю как доступная операция.", body))
    story.append(Paragraph("Что просим у Octo", h))
    for x in [
        "Подтвердить разрешённые P2P/payout/merchant сценарии и юридическую роль B1Pay.",
        "Предоставить sandbox/API/signing/webhook/error-code documentation.",
        "Согласовать лимиты, комиссии, settlement, refund/reversal/chargeback и SLA.",
        "Определить AML/KYC RACI, security requirements, пилотную группу и KPI.",
    ]: story.append(Paragraph("• " + x, bullet))
    story.append(Spacer(1, 4))
    story.append(Paragraph("Юридическое ограничение", h))
    story.append(Paragraph("B1Pay не заявляет себя банком, платёжной организацией, МФО, страховщиком или инвестиционным посредником. Этот документ — концепция для переговоров, не юридическое заключение. До запуска требуется консультация юриста и письменное подтверждение Octo/партнёров.", body))
    story.append(Spacer(1, 4))
    story.append(Paragraph("Приложения: business_model.md • architecture.md • platform_architecture.svg • полная презентация PPTX/PDF", small))
    doc.build(story)
    return ONE_PAGER_PATH


if __name__ == "__main__":
    print(build_pptx())
    print(build_one_pager())
    # Convert PPTX to PDF with LibreOffice where available.
    soffice = "/Users/macbookpro/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"
    if os.path.exists(soffice):
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(PPTX_PATH)], check=False)
        generated = OUT / (PPTX_PATH.stem + ".pdf")
        if generated.exists():
            print(generated)
