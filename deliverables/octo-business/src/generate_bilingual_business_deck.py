from __future__ import annotations

import os
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
PPTX_PATH = OUT / "B1Pay_Bilingual_Business_Partner_Proposal.pptx"
PDF_PATH = OUT / "B1Pay_Bilingual_Business_Partner_Proposal.pdf"

W = Inches(13.333)
H = Inches(7.5)
NAVY = RGBColor(17, 42, 77)
BLUE = RGBColor(18, 100, 214)
CYAN = RGBColor(44, 177, 213)
INK = RGBColor(32, 45, 62)
MUTED = RGBColor(86, 108, 135)
PALE = RGBColor(242, 247, 252)
PALE_BLUE = RGBColor(233, 242, 255)
PALE_GREEN = RGBColor(237, 248, 241)
PALE_ORANGE = RGBColor(255, 244, 223)
PALE_RED = RGBColor(255, 241, 237)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(183, 203, 226)
GREEN = RGBColor(36, 150, 95)
ORANGE = RGBColor(245, 181, 71)
RED = RGBColor(194, 85, 72)


def add_full_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element); slide.shapes._spTree.insert(2, shape._element)


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP, margin=0.04, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text; run.font.name = "Aptos"; run.font.size = Pt(size)
    run.font.bold = bold; run.font.italic = italic; run.font.color.rgb = color
    return tb


def add_footer(slide, idx):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(7.12), Inches(12.38), Inches(0.008))
    line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background()
    add_text(slide, 0.52, 7.18, 8.5, 0.18, "B1Pay × Octo | Русский / O‘zbekcha", 8, MUTED)
    add_text(slide, 12.35, 7.16, 0.45, 0.2, str(idx), 9, MUTED, align=PP_ALIGN.RIGHT)


def add_brand(slide, dark=False):
    color = WHITE if dark else NAVY
    add_text(slide, 0.55, 0.18, 2.0, 0.33, "B1Pay", 18, color, True)


def bilingual(slide, x, y, w, ru, uz, size=18, color=INK, bold=False, uz_size=None, h=0.8, align=PP_ALIGN.LEFT):
    uz_size = uz_size or max(8, int(size * 0.62))
    add_text(slide, x, y, w, h * 0.60, ru, size, color, bold, align=align)
    add_text(slide, x, y + h * 0.57, w, h * 0.40, "UZ: " + uz, uz_size, MUTED, False, align=align)


def title(slide, idx, ru, uz, subtitle_ru=None, subtitle_uz=None, dark=False):
    if dark:
        add_brand(slide, dark=True)
        add_text(slide, 0.72, 0.66, 11.6, 0.52, ru, 27, WHITE, True)
        add_text(slide, 0.76, 1.22, 11.6, 0.36, "UZ: " + uz, 13, RGBColor(190, 211, 235))
        if subtitle_ru:
            add_text(slide, 0.76, 1.75, 11.6, 0.45, subtitle_ru, 15, RGBColor(214, 229, 246))
            if subtitle_uz:
                add_text(slide, 0.76, 2.12, 11.6, 0.30, "UZ: " + subtitle_uz, 10, RGBColor(176, 198, 222))
    else:
        add_brand(slide)
        add_text(slide, 0.52, 0.54, 12.15, 0.48, ru, 25, NAVY, True)
        add_text(slide, 0.55, 1.03, 12.0, 0.30, "UZ: " + uz, 11, BLUE, True)
        if subtitle_ru:
            add_text(slide, 0.55, 1.37, 11.9, 0.30, subtitle_ru, 11, MUTED)
            if subtitle_uz:
                add_text(slide, 0.55, 1.64, 11.9, 0.26, "UZ: " + subtitle_uz, 9, MUTED)
        add_footer(slide, idx)


def box(slide, x, y, w, h, ru_title, uz_title, ru_body, uz_body, fill=PALE_BLUE, accent=BLUE, title_size=14, body_size=10.5):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.color.rgb = LINE; sh.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    add_text(slide, x + 0.18, y + 0.12, w - 0.30, 0.25, ru_title, title_size, NAVY, True)
    add_text(slide, x + 0.18, y + 0.40, w - 0.30, 0.20, "UZ: " + uz_title, max(8, int(title_size * 0.64)), BLUE, True)
    body = "RU: " + ru_body + "\n\nUZ: " + uz_body
    add_text(slide, x + 0.18, y + 0.72, w - 0.32, h - 0.82, body, body_size, INK)
    return sh


def bullet_box(slide, x, y, w, h, ru_title, uz_title, bullets, fill=PALE_BLUE, accent=BLUE, body_size=10.5):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.color.rgb = LINE; sh.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    add_text(slide, x + 0.18, y + 0.12, w - 0.3, 0.25, ru_title, 14, NAVY, True)
    add_text(slide, x + 0.18, y + 0.40, w - 0.3, 0.2, "UZ: " + uz_title, 9, BLUE, True)
    text = ""
    for i, (ru, uz) in enumerate(bullets):
        if i: text += "\n"
        text += "• " + ru + "\n  UZ: " + uz
    add_text(slide, x + 0.18, y + 0.75, w - 0.34, h - 0.84, text, body_size, INK)
    return sh


def connector(slide, x1, y1, x2, y2, color=BLUE, width=2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    return c


def table(slide, x, y, w, h, headers, rows, widths=None, font_size=9.5):
    cols = len(headers); nrows = len(rows) + 1
    shape = slide.shapes.add_table(nrows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if widths:
        for i, width in enumerate(widths): tbl.columns[i].width = Inches(width)
    for r in range(nrows):
        for c in range(cols):
            cell = tbl.cell(r, c); cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if r == 0 else (PALE_BLUE if r % 2 else PALE)
            cell.text = headers[c] if r == 0 else rows[r - 1][c]
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = "Aptos"; run.font.size = Pt(font_size if r else font_size + 0.3)
                    run.font.bold = r == 0; run.font.color.rgb = WHITE if r == 0 else INK
    return shape


def new_slide(prs, idx, ru, uz, subtitle_ru=None, subtitle_uz=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_full_bg(slide, NAVY if dark else WHITE)
    title(slide, idx, ru, uz, subtitle_ru, subtitle_uz, dark)
    return slide


def build():
    prs = Presentation(); prs.slide_width = W; prs.slide_height = H
    prs.core_properties.title = "B1Pay × Octo — двуязычная бизнес-модель и предложение о партнёрстве"
    prs.core_properties.subject = "Бизнес-модель, финансы, пилот, роли и интеграция / Biznes modeli, moliya, pilot, rollar va integratsiya"
    prs.core_properties.author = "B1Pay"
    prs.core_properties.keywords = "B1Pay, Octo, Uzbekistan, B-PAY, MyID, P2P, fintech"

    # 1 cover
    s = new_slide(prs, 1, "B1Pay × Octo", "B1Pay × Octo", "Бизнес-модель и предложение о партнёрстве для Узбекистана", "O‘zbekiston uchun biznes modeli va hamkorlik taklifi", dark=True)
    add_text(s, 0.76, 2.62, 11.6, 0.58, "Сравнить → выбрать → подтвердить → получить сервис через лицензированного партнёра", 25, WHITE, True)
    add_text(s, 0.76, 3.20, 11.6, 0.32, "UZ: Solishtirish → tanlash → tasdiqlash → litsenziyalangan hamkor orqali xizmat olish", 11, RGBColor(190, 211, 235))
    box(s, 0.76, 4.35, 3.6, 1.30, "B1Pay", "B1Pay", "Сайт, приложение, каталог, заявки, MyID/QR, аналитика", "Sayt, ilova, katalog, arizalar, MyID/QR va tahlil", fill=RGBColor(31, 62, 104), accent=CYAN, title_size=15, body_size=10)
    box(s, 4.87, 4.35, 3.6, 1.30, "Octo / банк", "Octo / bank", "Платёжная инфраструктура и регулируемая операция в согласованной роли", "To‘lov infratuzilmasi va kelishilgan roldagi tartibga solinadigan operatsiya", fill=RGBColor(31, 62, 104), accent=BLUE, title_size=15, body_size=9.5)
    box(s, 8.98, 4.35, 3.6, 1.30, "Пользователь", "Foydalanuvchi", "Понятные условия, подтверждение личности и статус заявки", "Tushunarli shartlar, shaxsni tasdiqlash va ariza holati", fill=RGBColor(31, 62, 104), accent=ORANGE, title_size=15, body_size=10)
    add_text(s, 0.78, 6.65, 11.3, 0.28, "Концепция для переговоров • 10 сентября 2026", 10, RGBColor(176, 198, 222))
    add_text(s, 0.78, 6.91, 11.3, 0.22, "Muzokaralar uchun konsepsiya • 2026-yil 10-sentabr", 8, RGBColor(176, 198, 222))

    # 2 summary
    s = new_slide(prs, 2, "Краткое резюме", "Qisqacha xulosa", "Что именно предлагаем партнёру", "Hamkorga nimani taklif qilamiz")
    bullet_box(s, 0.65, 1.95, 5.85, 3.75, "Проблема рынка", "Bozor muammosi", [
        ("Пользователь сравнивает предложения на разных сайтах.", "Foydalanuvchi takliflarni turli saytlarda solishtiradi."),
        ("Условия, комиссии и требования часто непонятны.", "Shartlar, komissiyalar va talablar ko‘pincha tushunarsiz."),
        ("Заявка и подтверждение личности проходят в разных системах.", "Ariza va shaxsni tasdiqlash turli tizimlarda amalga oshadi."),
    ], fill=PALE_RED, accent=RED, body_size=11)
    bullet_box(s, 6.8, 1.95, 5.85, 3.75, "Решение B1Pay", "B1Pay yechimi", [
        ("Одна витрина банковских и финансовых предложений.", "Bank va moliyaviy takliflar uchun yagona vitrina."),
        ("Сравнение условий, калькуляторы и понятная заявка.", "Shartlarni solishtirish, kalkulyatorlar va tushunarli ariza."),
        ("MyID на телефоне и QR-вход на сайте.", "Telefonda MyID va saytda QR orqali kirish."),
        ("Операции выполняются лицензированным партнёром.", "Operatsiyalar litsenziyalangan hamkor tomonidan bajariladi."),
    ], fill=PALE_GREEN, accent=GREEN, body_size=10.7)
    add_text(s, 0.8, 6.10, 11.5, 0.46, "Цель пилота: проверить спрос, качество заявок и безопасное техническое взаимодействие до расширения услуг.", 14, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.8, 6.48, 11.5, 0.28, "Pilot maqsadi: xizmatlarni kengaytirishdan oldin talab, ariza sifati va xavfsiz texnik hamkorlikni tekshirish.", 9, BLUE, align=PP_ALIGN.CENTER)

    # 3 problem/solution flow
    s = new_slide(prs, 3, "Путь пользователя", "Foydalanuvchi yo‘li", "От поиска продукта до решения партнёра", "Mahsulotni izlashdan hamkor qarorigacha")
    stages = [
        ("1. Найти", "1. Topish", "Каталог банков и продуктов", "Banklar va mahsulotlar katalogi", PALE_BLUE, BLUE),
        ("2. Сравнить", "2. Solishtirish", "Ставка, срок, комиссия, требования", "Foiz, muddat, komissiya va talablar", PALE_ORANGE, ORANGE),
        ("3. Подтвердить", "3. Tasdiqlash", "MyID и согласие пользователя", "MyID va foydalanuvchi roziligi", PALE_GREEN, GREEN),
        ("4. Подать", "4. Yuborish", "Заявка у выбранного партнёра", "Tanlangan hamkorga ariza", PALE_BLUE, CYAN),
        ("5. Получить статус", "5. Holatni olish", "Решение банка/партнёра", "Bank/hamkor qarori", PALE_RED, RED),
    ]
    xs = [0.55, 3.10, 5.65, 8.20, 10.75]
    for i, (rt, ut, rb, ub, fill, accent) in enumerate(stages):
        box(s, xs[i], 2.35, 2.15, 2.65, rt, ut, rb, ub, fill=fill, accent=accent, title_size=13, body_size=9.7)
        if i < 4: connector(s, xs[i] + 2.17, 3.65, xs[i + 1] - 0.05, 3.65)
    add_text(s, 0.75, 5.55, 11.7, 0.44, "B1Pay не принимает решение по кредиту или страховке: окончательное решение остаётся у лицензированного поставщика.", 13, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.75, 5.95, 11.7, 0.30, "B1Pay kredit yoki sug‘urta bo‘yicha qaror qabul qilmaydi: yakuniy qaror litsenziyalangan provayderda qoladi.", 9, BLUE, align=PP_ALIGN.CENTER)

    # 4 product scope table
    s = new_slide(prs, 4, "Продукты и границы запуска", "Mahsulotlar va ishga tushirish chegaralari", "Что доступно сейчас и что требует отдельного согласования", "Hozir nima mavjud va nimaga alohida kelishuv kerak")
    headers = ["Направление / Yo‘nalish", "Сейчас / Hozir", "После согласования / Kelishuvdan keyin"]
    rows = [
        ["Кредиты и микрозаймы\nKredit va mikroqarz", "Сравнение, калькулятор, заявка\nSolishtirish, kalkulyator, ariza", "Решение и договор у банка/МФО\nQaror va shartnoma bank/MFOda"],
        ["Карты и депозиты\nKartalar va depozitlar", "Каталог условий и заявка\nShartlar katalogi va ariza", "Открытие/обслуживание через партнёра\nHamkor orqali ochish/xizmat"],
        ["Страхование\nSug‘urta", "Сравнение полисов и заявка\nPolislarni solishtirish va ariza", "Полис и claims у страховщика\nPolis va da’volar sug‘urtachida"],
        ["Инвестиционный модуль\nInvestitsiya moduli", "Каталог с раскрытием рисков\nXavflar ko‘rsatilgan katalog", "Отдельная правовая квалификация\nAlohida huquqiy baholash"],
        ["P2P / переводы\nP2P / o‘tkazmalar", "Выключено до approval gate\nTasdiqlashgacha o‘chirilgan", "Только через Octo/банк и разрешённую модель\nFaqat Octo/bank va ruxsat etilgan model orqali"],
    ]
    table(s, 0.55, 2.05, 12.25, 4.45, headers, rows, widths=[2.35, 4.45, 5.45], font_size=8.6)

    # 5 business model
    s = new_slide(prs, 5, "Бизнес-модель B1Pay", "B1Pay biznes modeli", "B1Pay — технологическая платформа и канал привлечения клиентов", "B1Pay — texnologik platforma va mijozlarni jalb qilish kanali")
    box(s, 0.70, 2.00, 3.55, 2.60, "Пользователь", "Foydalanuvchi", "Выбирает продукт, сравнивает условия и даёт согласие на передачу заявки.", "Mahsulotni tanlaydi, shartlarni solishtiradi va arizani yuborishga rozilik beradi.", fill=PALE_BLUE, accent=BLUE, body_size=10.2)
    box(s, 4.90, 2.00, 3.55, 2.60, "B1Pay", "B1Pay", "Показывает каталог, собирает заявку, подтверждает согласие и направляет клиента партнёру.", "Katalogni ko‘rsatadi, arizani yig‘adi, rozilikni tasdiqlaydi va mijozni hamkorga yo‘naltiradi.", fill=PALE_GREEN, accent=GREEN, body_size=9.8)
    box(s, 9.10, 2.00, 3.55, 2.60, "Банк / партнёр", "Bank / hamkor", "Проверяет клиента, принимает решение, заключает договор и предоставляет услугу.", "Mijozni tekshiradi, qaror qabul qiladi, shartnoma tuzadi va xizmat ko‘rsatadi.", fill=PALE_ORANGE, accent=ORANGE, body_size=9.8)
    connector(s, 4.30, 3.30, 4.85, 3.30); connector(s, 8.50, 3.30, 9.05, 3.30)
    bullet_box(s, 1.0, 5.15, 11.25, 1.05, "Главный принцип", "Asosiy tamoyil", [
        ("B1Pay не является банком, МФО, страховщиком или платёжной организацией без отдельного правового основания.", "B1Pay alohida huquqiy asos bo‘lmasa bank, MFO, sug‘urta yoki to‘lov tashkiloti hisoblanmaydi."),
    ], fill=PALE_RED, accent=RED, body_size=10.5)

    # 6 revenue model
    s = new_slide(prs, 6, "Модель доходов", "Daromad modeli", "Прозрачная монетизация без скрытого влияния на рейтинг", "Reytingga yashirin ta’sirsiz shaffof monetizatsiya")
    headers = ["Источник / Manba", "Как работает / Qanday ishlaydi", "Условие / Shart"]
    rows = [
        ["CPA / Referral", "Плата за одобренную заявку или клиента\nTasdiqlangan ariza yoki mijoz uchun to‘lov", "Только по договору с партнёром\nFaqat hamkor shartnomasi asosida"],
        ["Lead fee", "Оплата за квалифицированную заявку\nSifatli ariza uchun to‘lov", "Критерии качества заранее\nSifat mezonlari oldindan"],
        ["Реклама / Reklama", "Платное размещение предложения\nTaklifni pullik joylashtirish", "Явная маркировка sponsored\nSponsored belgisi ko‘rsatiladi"],
        ["B2B / API", "Кабинет и интеграция для партнёров\nHamkorlar uchun kabinet va integratsiya", "SLA и техническая поддержка\nSLA va texnik yordam"],
        ["Premium services", "Дополнительные аналитические сервисы\nQo‘shimcha analitik xizmatlar", "Без обещаний гарантированного дохода\nKafolatlangan daromad va’dasisiz"],
    ]
    table(s, 0.55, 1.95, 12.25, 3.75, headers, rows, widths=[2.1, 6.0, 4.15], font_size=8.6)
    add_text(s, 0.9, 5.95, 11.4, 0.34, "Воронка: просмотр → заявка → проверка личности → решение → операция → доход.", 14, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.9, 6.30, 11.4, 0.27, "Voronka: ko‘rish → ariza → shaxsni tekshirish → qaror → operatsiya → daromad.", 9, BLUE, align=PP_ALIGN.CENTER)

    # 7 financial scenario
    s = new_slide(prs, 7, "Финансовый сценарий", "Moliyaviy ssenariy", "Базовые ориентиры для обсуждения, а не обещание результата", "Muhokama uchun bazaviy ko‘rsatkichlar, natija va’dasi emas")
    headers = ["Показатель / Ko‘rsatkich", "Базовый сценарий / Bazaviy ssenariy", "Комментарий / Izoh"]
    rows = [
        ["Стартовый капитал\nBoshlang‘ich kapital", "1 млрд сум\n1 mlrd so‘m", "Продукт, интеграции, команда и маркетинг\nMahsulot, integratsiya, jamoa va marketing"],
        ["Выручка за 3 года\n3 yillik daromad", "20,4 млрд сум\n20,4 mlrd so‘m", "Сценарий требует проверки фактическими данными\nAmaliy ma’lumotlar bilan tekshiriladi"],
        ["EBITDA за 3 года\n3 yillik EBITDA", "5,5 млрд сум\n5,5 mlrd so‘m", "Зависит от конверсии и комиссии партнёров\nKonversiya va hamkor komissiyasiga bog‘liq"],
        ["Цель 3-го года\n3-yil maqsadi", "14 млрд сум выручки / 4 млрд EBITDA\n14 mlrd daromad / 4 mlrd EBITDA", "Целевой ориентир, не гарантия\nMaqsadli ko‘rsatkich, kafolat emas"],
    ]
    table(s, 0.55, 2.00, 12.25, 3.55, headers, rows, widths=[2.6, 4.0, 5.65], font_size=8.8)
    bullet_box(s, 1.0, 5.62, 11.25, 1.18, "Важно", "Muhim", [
        ("Все цифры — предварительный базовый сценарий для финансового обсуждения; после пилота их нужно обновить фактическими KPI.", "Barcha raqamlar moliyaviy muhokama uchun dastlabki bazaviy ssenariy; pilotdan so‘ng haqiqiy KPI bilan yangilanadi."),
    ], fill=PALE_ORANGE, accent=ORANGE, body_size=9.3)

    # 8 unit economics
    s = new_slide(prs, 8, "Unit economics", "Unit economics", "Экономика одного привлечённого клиента", "Bitta jalb qilingan mijoz iqtisodiyoti")
    headers = ["Метрика / Metrika", "Ориентир / Ko‘rsatkich", "Что означает / Ma’nosi"]
    rows = [
        ["CAC", "≈ 20 000 сум", "Стоимость привлечения клиента\nMijozni jalb qilish qiymati"],
        ["ARPU", "≈ 60 000 сум", "Средний доход на клиента\nBitta mijozdan o‘rtacha daromad"],
        ["LTV", "150 000–200 000 сум", "Доход за весь срок взаимодействия\nButun munosabat davridagi daromad"],
        ["LTV / CAC", "7,5–10×", "Запас для масштабирования модели\nModelni kengaytirish uchun zaxira"],
    ]
    table(s, 0.85, 2.0, 7.3, 3.65, headers, rows, widths=[1.75, 1.85, 3.70], font_size=9)
    box(s, 8.55, 2.0, 3.95, 3.65, "Цель / Maqsad", "Maqsad", "Стоимость привлечения должна быть ниже дохода от клиента за весь срок. Показатели нужно подтвердить после запуска рекламных каналов и пилота.", "Mijozni jalb qilish qiymati uning butun davrdagi daromadidan past bo‘lishi kerak. Ko‘rsatkichlar reklama kanallari va pilotdan keyin tasdiqlanadi.", fill=PALE_GREEN, accent=GREEN, title_size=16, body_size=11)

    # 9 break even
    s = new_slide(prs, 9, "Точка операционной безубыточности", "Operatsion zararsizlik nuqtasi", "Пример расчёта на базовом сценарии", "Bazaviy ssenariy bo‘yicha hisob-kitob")
    box(s, 0.8, 2.1, 3.45, 2.25, "Постоянные расходы", "Doimiy xarajatlar", "300 млн сум в месяц", "Oyiga 300 mln so‘m", fill=PALE_BLUE, accent=BLUE, title_size=15, body_size=14)
    box(s, 4.95, 2.1, 3.45, 2.25, "Доход с клиента", "Bitta mijoz daromadi", "60 000 сум contribution", "60 000 so‘m contribution", fill=PALE_ORANGE, accent=ORANGE, title_size=15, body_size=13)
    box(s, 9.10, 2.1, 3.45, 2.25, "Break-even", "Zararsizlik", "300 000 000 ÷ 60 000 = 5 000 подтверждённых клиентов в месяц", "300 000 000 ÷ 60 000 = oyiga 5 000 ta tasdiqlangan mijoz", fill=PALE_GREEN, accent=GREEN, title_size=15, body_size=11)
    connector(s, 4.28, 3.22, 4.88, 3.22); connector(s, 8.43, 3.22, 9.03, 3.22)
    add_text(s, 0.85, 5.15, 11.6, 0.50, "Это ориентир, а не финансовая гарантия: фактическая точка зависит от расходов, комиссии и конверсии.", 14, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.85, 5.58, 11.6, 0.30, "Bu yo‘nalish, moliyaviy kafolat emas: haqiqiy nuqta xarajatlar, komissiya va konversiyaga bog‘liq.", 9, BLUE, align=PP_ALIGN.CENTER)

    # 10 partner value
    s = new_slide(prs, 10, "Ценность для Octo и банков", "Octo va banklar uchun qiymat", "Что получает партнёр от совместного пилота", "Hamkor qo‘shma pilotdan nimani oladi")
    bullet_box(s, 0.65, 1.95, 5.85, 3.95, "Для Octo / банка", "Octo / bank uchun", [
        ("Новый цифровой канал привлечения клиентов.", "Mijozlarni jalb qilishning yangi raqamli kanali."),
        ("Структурированные заявки и единая точка интеграции.", "Tuzilgan arizalar va yagona integratsiya nuqtasi."),
        ("Контролируемый тест: лимиты, мониторинг и отчётность.", "Nazorat qilinadigan test: limitlar, monitoring va hisobot."),
        ("Понятная воронка: просмотр → заявка → решение.", "Tushunarli voronka: ko‘rish → ariza → qaror."),
    ], fill=PALE_GREEN, accent=GREEN, body_size=10.8)
    bullet_box(s, 6.8, 1.95, 5.85, 3.95, "Что остаётся у B1Pay", "B1Payda qoladiganlar", [
        ("UX, каталог, сравнение и первая линия поддержки.", "UX, katalog, solishtirish va birinchi darajali yordam."),
        ("Согласие пользователя и маршрутизация заявки.", "Foydalanuvchi roziligi va arizani yo‘naltirish."),
        ("Аналитика качества трафика и конверсии.", "Trafik sifati va konversiya tahlili."),
        ("P2P включается только после отдельного approval gate.", "P2P faqat alohida tasdiqlash bosqichidan keyin yoqiladi."),
    ], fill=PALE_BLUE, accent=BLUE, body_size=10.8)

    # 11 roles table
    s = new_slide(prs, 11, "Роли и ответственность", "Rollar va mas’uliyat", "Кто за что отвечает в партнёрской модели", "Hamkorlik modelida kim nimaga javob beradi")
    headers = ["Участник / Ishtirokchi", "Ответственность RU / Mas’uliyat UZ", "Результат / Natija"]
    rows = [
        ["B1Pay", "UX, каталог, заявки, consent, поддержка L1\nUX, katalog, arizalar, rozilik, L1 yordam", "Качественный клиентский путь\nSifatli mijoz yo‘li"],
        ["Octo / банк", "API, лимиты, расчёты, fraud и платёжный статус\nAPI, limitlar, hisob-kitob, fraud va to‘lov holati", "Безопасная операция в разрешённой роли\nRuxsat etilgan roldagi xavfsiz operatsiya"],
        ["Банк / МФО", "Решение по кредиту, договор и выдача\nKredit qarori, shartnoma va berish", "Финансовый продукт клиенту\nMijozga moliyaviy mahsulot"],
        ["Страховщик", "Полис, андеррайтинг и claims\nPolis, anderrayting va da’volar", "Страховое покрытие\nSug‘urta qoplamasi"],
        ["MyID", "Подтверждение личности в согласованном контуре\nKelishilgan konturda shaxsni tasdiqlash", "Достоверная KYC-сессия\nIshonchli KYC sessiyasi"],
    ]
    table(s, 0.55, 1.95, 12.25, 4.65, headers, rows, widths=[1.85, 7.10, 3.30], font_size=8.1)

    # 12 MyID QR
    s = new_slide(prs, 12, "MyID + QR: сайт и телефон работают вместе", "MyID + QR: sayt va telefon birga ishlaydi", "Web показывает QR, подтверждение выполняется в приложении", "Sayt QR ko‘rsatadi, tasdiqlash ilovada bajariladi")
    box(s, 0.65, 2.0, 3.45, 3.55, "Сайт", "Sayt", "1. Создаёт короткую сессию\n2. Показывает QR-код\n3. Ждёт подтверждение", "1. Qisqa sessiya yaratadi\n2. QR kodni ko‘rsatadi\n3. Tasdiqlashni kutadi", fill=PALE_BLUE, accent=BLUE, body_size=11)
    box(s, 4.95, 2.0, 3.45, 3.55, "Приложение", "Ilova", "1. Вход через MyID\n2. Сканирование QR\n3. Подтверждение действия", "1. MyID orqali kirish\n2. QR ni skanerlash\n3. Harakatni tasdiqlash", fill=PALE_GREEN, accent=GREEN, body_size=11)
    box(s, 9.25, 2.0, 3.45, 3.55, "Сервер B1Pay", "B1Pay serveri", "TTL, nonce, защита от повторного использования, журнал событий и минимальный набор данных партнёру.", "TTL, nonce, qayta ishlatishdan himoya, voqealar jurnali va hamkorga minimal ma’lumot.", fill=PALE_ORANGE, accent=ORANGE, body_size=10.5)
    connector(s, 4.15, 3.75, 4.90, 3.75); connector(s, 8.45, 3.75, 9.20, 3.75)
    add_text(s, 0.8, 6.02, 11.7, 0.42, "Биометрия и фотографии документов не должны попадать в обычные логи B1Pay.", 13, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.8, 6.37, 11.7, 0.26, "Biometriya va hujjat fotosuratlari B1Payning oddiy loglariga tushmasligi kerak.", 9, BLUE, align=PP_ALIGN.CENTER)

    # 13 architecture
    s = new_slide(prs, 13, "Техническая архитектура", "Texnik arxitektura", "Единый контур для сайта, Android и партнёрских API", "Sayt, Android va hamkor API’lari uchun yagona kontur")
    box(s, 0.65, 1.85, 2.55, 1.05, "Сайт", "Sayt", "Каталог · QR · заявки", "Katalog · QR · arizalar", fill=PALE_BLUE, accent=BLUE, body_size=9.2)
    box(s, 3.45, 1.85, 2.55, 1.05, "Android", "Android", "MyID · QR · статусы", "MyID · QR · holatlar", fill=PALE_BLUE, accent=CYAN, body_size=9.2)
    box(s, 6.25, 1.85, 2.55, 1.05, "Админ-панель", "Admin panel", "Офферы · заявки · аудит", "Takliflar · arizalar · audit", fill=PALE_ORANGE, accent=ORANGE, body_size=9.2)
    box(s, 9.05, 1.85, 3.55, 1.05, "Партнёрский кабинет", "Hamkor kabineti", "Лиды · SLA · отчёты", "Lidlar · SLA · hisobotlar", fill=PALE_GREEN, accent=GREEN, body_size=9.2)
    box(s, 1.20, 3.45, 11.05, 1.35, "B1Pay API и шлюз интеграций", "B1Pay API va integratsiya shlyuzi", "Сессии, каталог, согласия, заявки, статусы, webhooks, сверка, риск и аудит.", "Sessiyalar, katalog, roziliklar, arizalar, holatlar, webhooks, solishtirish, xavf va audit.", fill=RGBColor(235, 242, 253), accent=BLUE, title_size=16, body_size=11)
    box(s, 0.65, 5.50, 2.55, 0.95, "MyID", "MyID", "KYC status", "KYC holati", fill=PALE_RED, accent=RED, body_size=9.5)
    box(s, 3.45, 5.50, 2.55, 0.95, "Octo / банк", "Octo / bank", "P2P · payout · settlement", "P2P · payout · settlement", fill=PALE_ORANGE, accent=ORANGE, body_size=9.2)
    box(s, 6.25, 5.50, 2.55, 0.95, "Банк / МФО", "Bank / MFO", "Кредитное решение", "Kredit qarori", fill=PALE_GREEN, accent=GREEN, body_size=9.5)
    box(s, 9.05, 5.50, 3.55, 0.95, "Страховщик / инвестиции", "Sug‘urtachi / investitsiya", "Отдельный правовой контур", "Alohida huquqiy kontur", fill=PALE_RED, accent=RED, body_size=9.2)
    for x in [1.92, 4.72, 7.52, 10.82]: connector(s, x, 2.92, x, 3.42, BLUE, 1.4)
    for x in [1.92, 4.72, 7.52, 10.82]: connector(s, x, 4.82, x, 5.47, BLUE, 1.4)

    # 14 P2P/legal
    s = new_slide(prs, 14, "P2P и переводы: только после согласования", "P2P va o‘tkazmalar: faqat kelishuvdan keyin", "Merchant/acquiring и P2P — разные рабочие направления", "Merchant/acquiring va P2P — turli ish yo‘nalishlari")
    steps = [
        ("1. Право", "1. Huquq", "Роль B1Pay, договор и лицензированный контур", "B1Pay roli, shartnoma va litsenziyalangan kontur", PALE_RED, RED),
        ("2. Техника", "2. Texnika", "Sandbox, подписи, webhooks, лимиты", "Sandbox, imzo, webhooks, limitlar", PALE_BLUE, BLUE),
        ("3. Контроль", "3. Nazorat", "AML/KYC, fraud, reconciliation, support", "AML/KYC, fraud, reconciliation, yordam", PALE_ORANGE, ORANGE),
        ("4. Пилот", "4. Pilot", "Ограниченная группа и feature flag", "Cheklangan guruh va feature flag", PALE_GREEN, GREEN),
    ]
    xs = [0.75, 3.90, 7.05, 10.20]
    for i, (rt, ut, rb, ub, fill, accent) in enumerate(steps):
        box(s, xs[i], 2.25, 2.45, 2.80, rt, ut, rb, ub, fill=fill, accent=accent, title_size=13, body_size=10)
        if i < 3: connector(s, xs[i] + 2.47, 3.65, xs[i + 1] - 0.05, 3.65)
    add_text(s, 0.75, 5.62, 11.7, 0.45, "До прохождения всех этапов P2P остаётся выключенным и не рекламируется как доступная услуга.", 14, RED, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.75, 6.02, 11.7, 0.27, "Barcha bosqichlar o‘tilmaguncha P2P o‘chiriladi va mavjud xizmat sifatida reklama qilinmaydi.", 9, RED, align=PP_ALIGN.CENTER)

    # 15 compliance
    s = new_slide(prs, 15, "Безопасность и защита клиента", "Xavfsizlik va mijozni himoya qilish", "Обязательные условия до коммерческого запуска", "Tijoriy ishga tushirishdan oldingi majburiy shartlar")
    bullet_box(s, 0.65, 1.95, 3.85, 4.10, "KYC / AML", "KYC / AML", [
        ("Ответственный партнёр принимает регулируемое решение.", "Mas’ul hamkor tartibga solinadigan qarorni qabul qiladi."),
        ("Письменное распределение ролей и эскалаций.", "Rollar va eskalatsiyalar yozma belgilanadi."),
        ("Журнал согласий и доказательств.", "Roziliklar va dalillar jurnali."),
    ], fill=PALE_BLUE, accent=BLUE, body_size=10.4)
    bullet_box(s, 4.75, 1.95, 3.85, 4.10, "Техническая защита", "Texnik himoya", [
        ("Секреты только на сервере; подпись запросов.", "Sirlar faqat serverda; so‘rovlarni imzolash."),
        ("Не хранить raw PAN, CVV и биометрию в логах.", "Raw PAN, CVV va biometrikani loglarda saqlamaslik."),
        ("Роли, MFA, мониторинг и план инцидентов.", "Rollar, MFA, monitoring va hodisa rejasi."),
    ], fill=PALE_GREEN, accent=GREEN, body_size=10.4)
    bullet_box(s, 8.85, 1.95, 3.80, 4.10, "Защита потребителя", "Iste’molchini himoya qilish", [
        ("Источник и дата обновления условий.", "Manba va shartlarning yangilanish sanasi."),
        ("Понятные комиссии, риски и жалобы.", "Tushunarli komissiyalar, xavflar va shikoyatlar."),
        ("Нет обещаний гарантированного дохода.", "Kafolatlangan daromad va’dasi yo‘q."),
    ], fill=PALE_RED, accent=RED, body_size=10.4)

    # 16 roadmap
    s = new_slide(prs, 16, "Дорожная карта", "Yo‘l xaritasi", "Порядок запуска от legal discovery до расширения", "Huquqiy kelishuvdan kengaytirishgacha ishga tushirish tartibi")
    roadmap = [
        ("0", "Право и discovery", "Huquq va discovery", "Роль, договор, RACI, sandbox", "Rol, shartnoma, RACI, sandbox", PALE_RED, RED),
        ("1", "Каталог и заявки", "Katalog va arizalar", "Офферы, consent, MyID/QR", "Takliflar, rozilik, MyID/QR", PALE_BLUE, BLUE),
        ("2", "Ограниченный пилот", "Cheklangan pilot", "1–2 партнёра, лимиты, KPI", "1–2 hamkor, limitlar, KPI", PALE_GREEN, GREEN),
        ("3", "Расширение", "Kengaytirish", "P2P/payout после approval", "Tasdiqlashdan keyin P2P/payout", PALE_ORANGE, ORANGE),
    ]
    for i, (n, rt, ut, rb, ub, fill, accent) in enumerate(roadmap):
        x = 0.60 + i * 3.16
        box(s, x, 2.10, 2.72, 2.85, f"{n}. {rt}", ut, rb, ub, fill=fill, accent=accent, title_size=12.7, body_size=10.2)
        if i < 3: connector(s, x + 2.75, 3.50, x + 3.10, 3.50)
    add_text(s, 0.80, 5.65, 11.5, 0.40, "P2P включается только после юридического, договорного и технического approval gate.", 14, RED, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.80, 6.00, 11.5, 0.27, "P2P faqat huquqiy, shartnomaviy va texnik tasdiqlash bosqichidan keyin yoqiladi.", 9, RED, align=PP_ALIGN.CENTER)

    # 17 partner ask
    s = new_slide(prs, 17, "Что нужно согласовать с Octo", "Octo bilan nimani kelishish kerak", "Конкретные решения для начала пилота", "Pilotni boshlash uchun aniq qarorlar")
    bullet_box(s, 0.65, 1.95, 5.85, 4.20, "Коммерция и право", "Tijorat va huquq", [
        ("Допустимая роль B1Pay и список услуг.", "B1Payning ruxsat etilgan roli va xizmatlar ro‘yxati."),
        ("Кто заключает договор с пользователем.", "Foydalanuvchi bilan shartnomani kim tuzadi."),
        ("Комиссии, расчёты, возвраты и споры.", "Komissiyalar, hisob-kitob, qaytarish va nizolar."),
        ("AML/KYC и ответственность за отчётность.", "AML/KYC va hisobot uchun mas’uliyat."),
    ], fill=PALE_ORANGE, accent=ORANGE, body_size=10.7)
    bullet_box(s, 6.80, 1.95, 5.85, 4.20, "Техника и пилот", "Texnika va pilot", [
        ("Тестовая среда, API, подписи и webhooks.", "Test muhiti, API, imzolar va webhooks."),
        ("Лимиты, idempotency, сверка и тестовые данные.", "Limitlar, idempotency, solishtirish va test ma’lumotlari."),
        ("Security requirements и поддержка инцидентов.", "Xavfsizlik talablari va hodisalarni qo‘llab-quvvatlash."),
        ("Группа пилота, срок, KPI и go/no-go.", "Pilot guruhi, muddat, KPI va go/no-go."),
    ], fill=PALE_BLUE, accent=BLUE, body_size=10.7)

    # 18 KPIs
    s = new_slide(prs, 18, "Показатели успеха пилота", "Pilot muvaffaqiyati ko‘rsatkichlari", "Пилот оценивается не только по числу операций", "Pilot faqat operatsiyalar soni bilan baholanmaydi")
    headers = ["Направление / Yo‘nalish", "Метрики RU / UZ", "Зачем / Nima uchun"]
    rows = [
        ["Воронка / Voronka", "Просмотр → заявка → KYC → решение\nKo‘rish → ariza → KYC → qaror", "Понимать конверсию\nKonversiyani tushunish"],
        ["Надёжность / Ishonchlilik", "Success rate, pending time, webhook delivery\nSuccess rate, kutish va webhook", "Контролировать сервис\nXizmatni nazorat qilish"],
        ["Риск / Xavf", "Fraud, AML cases, complaints, disputes\nFraud, AML, shikoyatlar, nizolar", "Защищать клиента и партнёра\nMijoz va hamkorni himoya qilish"],
        ["Экономика / Iqtisodiyot", "CAC, CPA/CPL, revenue per lead\nCAC, CPA/CPL, bitta lid daromadi", "Понимать окупаемость\nQaytish muddatini tushunish"],
    ]
    table(s, 0.55, 2.05, 12.25, 3.75, headers, rows, widths=[2.25, 6.15, 3.85], font_size=8.7)
    add_text(s, 0.8, 6.12, 11.5, 0.38, "Главная метрика: клиент получает прозрачный путь, партнёр — качественную заявку в контролируемом контуре.", 13, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.8, 6.45, 11.5, 0.24, "Asosiy ko‘rsatkich: mijoz shaffof yo‘lni, hamkor esa nazorat qilinadigan konturda sifatli arizani oladi.", 9, BLUE, align=PP_ALIGN.CENTER)

    # 19 investment/ask
    s = new_slide(prs, 19, "Инвестиционная возможность", "Investitsiya imkoniyati", "Для масштабирования после подтверждения product-market fit", "Product-market fit tasdiqlangandan keyin kengaytirish uchun")
    box(s, 0.75, 2.00, 3.65, 2.75, "На что нужны средства", "Mablag‘ nimaga kerak", "Команда, интеграции с 3–5 банками, каталог, AI-подбор, маркетинг и контроль качества.", "Jamoa, 3–5 bank bilan integratsiya, katalog, AI-moslashtirish, marketing va sifat nazorati.", fill=PALE_BLUE, accent=BLUE, body_size=10.4)
    box(s, 4.85, 2.00, 3.65, 2.75, "Что проверяем сначала", "Avval nimani tekshiramiz", "Спрос, конверсия, стоимость заявки, качество партнёрских решений и повторное использование.", "Talab, konversiya, ariza qiymati, hamkor qarorlari sifati va qayta foydalanish.", fill=PALE_GREEN, accent=GREEN, body_size=10.4)
    box(s, 8.95, 2.00, 3.65, 2.75, "Ограничение", "Cheklov", "Финансовые показатели являются базовым сценарием. Инвестиция не означает гарантированный доход.", "Moliyaviy ko‘rsatkichlar bazaviy ssenariy. Investitsiya kafolatlangan daromadni anglatmaydi.", fill=PALE_RED, accent=RED, body_size=10.4)
    add_text(s, 0.85, 5.45, 11.5, 0.52, "B1Pay делает выбор финансового продукта проще, быстрее и прозрачнее.", 20, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.85, 5.92, 11.5, 0.32, "B1Pay moliyaviy mahsulot tanlashni oddiy, tez va shaffof qiladi.", 11, BLUE, align=PP_ALIGN.CENTER)

    # 20 closing
    s = new_slide(prs, 20, "Следующий шаг", "Keyingi qadam", "Совместно определить законный и безопасный периметр запуска", "Ishga tushirishning qonuniy va xavfsiz doirasini birgalikda belgilash")
    add_text(s, 0.95, 2.00, 11.4, 0.42, "Предлагаем провести рабочую встречу и подтвердить:", 21, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.95, 2.36, 11.4, 0.28, "Ishchi uchrashuv o‘tkazib, quyidagilarni tasdiqlashni taklif qilamiz:", 11, BLUE, align=PP_ALIGN.CENTER)
    bullet_box(s, 1.95, 2.88, 9.45, 2.42, "Решения для старта", "Boshlash uchun qarorlar", [
        ("роль B1Pay и список разрешённых услуг", "B1Pay roli va ruxsat etilgan xizmatlar ro‘yxati"),
        ("тестовое подключение и требования безопасности", "test ulanishi va xavfsizlik talablari"),
        ("пилотный продукт, группа, лимиты и KPI", "pilot mahsulot, guruh, limitlar va KPI"),
        ("условия, при которых можно включить переводы", "o‘tkazmalarni yoqish shartlari"),
    ], fill=PALE_BLUE, accent=BLUE, body_size=12.0)
    add_text(s, 1.0, 5.75, 11.3, 0.46, "Спасибо. Контакт для рабочей встречи: команда B1Pay", 18, BLUE, True, align=PP_ALIGN.CENTER)
    add_text(s, 1.0, 6.15, 11.3, 0.28, "Rahmat. Ishchi uchrashuv uchun aloqa: B1Pay jamoasi", 10, MUTED, align=PP_ALIGN.CENTER)

    prs.save(PPTX_PATH)
    return PPTX_PATH


if __name__ == "__main__":
    print(build())
    soffice = "/Users/macbookpro/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"
    if os.path.exists(soffice):
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(PPTX_PATH)], check=False)
        generated = OUT / (PPTX_PATH.stem + ".pdf")
        if generated.exists():
            print(generated)
